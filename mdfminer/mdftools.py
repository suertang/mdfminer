# mdf.py 
# (C) 2017 Patrick Menschel

def to_csv_file(mdf_obj,fname,useabsolutetime=False,csv_sep=",",line_sep=";\n"):
    with open(fname,'w') as f:
        chans = mdf_obj.get_channel_short_names()
        f.write("time"+csv_sep+csv_sep.join(chans)+line_sep)
        records = mdf_obj.get_records_with_timestamp(useabsolutetime=useabsolutetime)
            
        for record in records:
            for timestamp in record:
                thisrecord = record[timestamp]
                f.write(str(timestamp)+csv_sep+csv_sep.join([str(thisrecord[chan]) for chan in chans])+line_sep)
    return
                        
def to_xlsx_file(mdf_obj,fname,useabsolutetime=False):
    from openpyxl import Workbook
    from openpyxl.chart import (
                                LineChart,
                                Reference,
                                )
    wb = Workbook()
    ws = wb.active
    ws.title = "data"
    row = ["time",]
    chans = mdf_obj.get_channel_short_names()
    col_idx = len(chans)+1
    row.extend(chans)
    ws.append(row)
    records = mdf_obj.get_records_with_timestamp(useabsolutetime=useabsolutetime)
    row_idx = 1
    for record in records:
            
        for timestamp in record:
                
            thisrecord = record[timestamp]
                
            row = [timestamp,]
            row.extend(thisrecord)
            ws.append(row)
            row_idx += 1
        
    c1 = LineChart()
    c1.title = "Line Chart"
    c1.style = 13
    c1.y_axis.title = 'Value'
    c1.x_axis.title = 'Record'
    
    data = Reference(ws, min_col=2, min_row=1, max_col=col_idx, max_row=row_idx)
    c1.add_data(data, titles_from_data=True)
    ws2 = wb.create_sheet("Chart")
    ws2.add_chart(c1, "A1")
    wb.save(fname)
    return
